"""DOCX text extraction (bonus format) via python-docx."""
from __future__ import annotations

from pathlib import Path

import docx


def read_docx(path: Path) -> str:
    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:  # tables often hold skills grids
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
